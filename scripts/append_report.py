# -*- coding: utf-8 -*-
"""Anexa seções finais ao relatório, preservando conteúdo + imagens existentes."""
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = "Relatório técnico H&W .docx"
OUT = "Relatório técnico H&W - completo.docx"

d = docx.Document(SRC)

def H(txt, lvl=3):
    d.add_paragraph(txt, style=f"Heading {lvl}")

def P(txt):
    d.add_paragraph(txt, style="normal")

def bullet(txt):
    try:
        d.add_paragraph(txt, style="List Bullet")
    except Exception:
        d.add_paragraph("• " + txt, style="normal")

def table(headers, rows):
    t = d.add_table(rows=1, cols=len(headers))
    try:
        t.style = "Table Grid"
    except Exception:
        pass
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    d.add_paragraph("")

# ---------------------------------------------------------------------------
H("Contenção & Hardening — mitigação por elo (antes/depois)", 3)
P("Para cada elo da cadeia de ataque foi aplicada uma mitigação e comprovada a "
  "mesma tentativa agora barrada. A filosofia foi defesa em camadas: corrigir tanto "
  "na rede quanto na aplicação, nunca depender de um único controle.")

H("Elo 2 — Contenção do movimento lateral (rede)", 4)
P("Antes, a regra de firewall F1 aceitava toda a sub-rede da DMZ (172.30.10.0/24) "
  "como origem para falar com a app-api — foi por isso que o DVWA comprometido "
  "(172.30.10.10) conseguiu pivotar para a camada APP. A mitigação restringiu F1 a "
  "aceitar apenas o WAF (172.30.10.5) como origem legítima.")
P("Resultado comprovado após a mudança:")
bullet("DVWA (172.30.10.10) → app-api:8000  →  TIMEOUT (bloqueado) — o movimento lateral morreu.")
bullet("WAF (172.30.10.5) → app-api:8000  →  succeeded — o caminho legítimo foi preservado.")
P("Impacto: uma única regra de rede interrompe a cadeia inteira a partir do Elo 2. O "
  "host de entrada comprometido deixa de ter rota para a camada de aplicação.")

H("Elo 3 — Vazamento de credencial (aplicação)", 4)
P("O endpoint interno /internal/db-config, que devolvia a credencial do banco sem "
  "autenticação, foi removido no modo endurecido (redução de superfície). Após a "
  "correção, a requisição retorna 404 — a credencial não é mais exposta.")

H("Elo 5 — RCE via command injection (aplicação)", 4)
P("O endpoint /internal/net-check concatenava a entrada diretamente no shell, "
  "permitindo execução de comando. A correção aplicou uma allowlist estrita de "
  "entrada (apenas caracteres válidos de IP/hostname) e passou a executar o comando "
  "sem shell (lista de argumentos, sem interpretação de metacaracteres). Após a "
  "correção: uma tentativa de injeção (host=127.0.0.1; id) retorna 400 (rejeitada), "
  "enquanto uma consulta legítima (host=127.0.0.1) continua funcionando.")

H("Redução de superfície — filtragem de egress", 4)
P("A política default-deny na cadeia FORWARD impede que os segmentos iniciem conexões "
  "para fora. Isso contém reverse shells e exfiltração para a internet. Comprovado: "
  "tanto app-api (APP) quanto DVWA (DMZ) recebem timeout ao tentar alcançar 8.8.8.8:53. "
  "Nenhum segmento fala com a internet — apenas os fluxos internos explícitos (F1/F2).")

H("Recomendação adicional — least-privilege no banco", 4)
P("O usuário do Postgres é superuser (padrão da imagem). Um hardening adicional "
  "recomendado é criar um role dedicado com permissão de SELECT apenas nas tabelas "
  "estritamente necessárias, sem superuser. Isso reduz o impacto de um SQLi ou de um "
  "reúso de credencial. Fica registrado como recomendação priorizada.")

# ---------------------------------------------------------------------------
H("Mapa de Mapeamento de Portas", 3)
P("Todas as portas dos serviços de back-end (dvwa, app-api, postgres) não são "
  "publicadas no host — apenas o gateway (VPN) e o acesso ao WAF via túnel são "
  "expostos. A tabela abaixo consolida o mapa após o hardening:")
table(
    ["Host", "Porta", "Proto", "Quem pode acessar", "Justificativa"],
    [
        ["wg-gateway (borda)", "51820", "UDP", "Qualquer (internet/host)", "Entrada da VPN (WireGuard) — único ponto público"],
        ["wg-gateway", "22", "TCP", "Sub-rede VPN (10.13.13.0/24)", "Admin SSH só via VPN; login apenas por chave"],
        ["waf", "8080", "TCP", "Via VPN (regra F-VPN)", "Ponto de entrada web; reverse proxy para o DVWA"],
        ["dvwa", "80", "TCP", "Somente o WAF (mesma DMZ)", "Alvo atrás do WAF; nunca acessado direto"],
        ["app-api", "8000", "TCP", "Somente o WAF (F1 endurecida)", "API interna; o DVWA não alcança mais"],
        ["postgres", "5432", "TCP", "Somente a camada APP (F2)", "Banco fechado; DMZ e internet nunca"],
    ],
)
P("Um caminho que deve ser sempre bloqueado, comprovadamente bloqueado: DVWA "
  "(172.30.10.10) → postgres:5432 resulta em timeout (recusado), enquanto o caminho "
  "legítimo app-api (172.30.20.10) → postgres:5432 é aberto. O timeout (estado "
  "'filtered' no nmap) é a assinatura do firewall descartando em silêncio — "
  "segmentação ativa, não uma porta meramente fechada.")

# ---------------------------------------------------------------------------
H("Resumo dos Achados (severidade e CVSS aproximado)", 3)
P("Foram exploradas nove vulnerabilidades reais — muito acima do mínimo de duas "
  "exigido. A tabela resume cada achado com a severidade estimada:")
table(
    ["#", "Vulnerabilidade", "Local", "Severidade (CVSS aprox.)"],
    [
        ["1", "SQL Injection (clássico)", "DVWA /sqli", "Crítico (9.8)"],
        ["2", "SQL Injection Blind (time-based)", "DVWA /sqli_blind", "Crítico (9.1)"],
        ["3", "XSS Refletido", "DVWA /xss_r", "Médio (6.1)"],
        ["4", "XSS DOM (via fragmento)", "DVWA /xss_d", "Médio (6.1)"],
        ["5", "Upload inseguro → Webshell (RCE)", "DVWA /upload", "Crítico (9.8)"],
        ["6", "Command Injection", "DVWA /exec", "Crítico (9.8)"],
        ["7", "CSRF (troca de senha)", "DVWA /csrf", "Alto (8.8)"],
        ["8", "Weak Session ID (previsível)", "DVWA /weak_id", "Alto (7.5)"],
        ["9", "Exposição de credencial + RCE no pivô", "app-api interno", "Crítico (9.8)"],
    ],
)
P("A cadeia de ponta a ponta encadeou os achados 5, 9 e o reúso da credencial: "
  "foothold na DMZ (webshell) → movimento lateral para a APP → execução de comando na "
  "app-api → uso da credencial roubada com o cliente psql → exfiltração do banco. É "
  "importante destacar que a segmentação de rede neutralizou o uso direto da "
  "credencial a partir da DMZ (o banco permaneceu inalcançável do segmento comprometido); "
  "a exfiltração só foi possível operando de dentro da camada APP, atravessando a "
  "cadeia legítima — o que demonstra que a defesa de rede funciona e que a brecha "
  "restante é de aplicação, endereçada no hardening.")

# ---------------------------------------------------------------------------
H("Conclusão, Trade-offs e Escala para a Nuvem Real", 3)
P("O laboratório demonstrou o ciclo completo de segurança — atacar, entender e "
  "conter — com evidência em cada etapa. A infraestrutura sobe com um único comando, "
  "de forma reprodutível, e nenhum segredo é versionado no repositório.")

H("Trade-offs e limitações", 4)
bullet("O WAF é um controle poderoso, mas não é bala de prata: o XSS DOM via fragmento "
       "demonstrou um ponto cego estrutural (o payload nunca chega ao servidor). A "
       "mitigação correta foi defesa em camadas — um cabeçalho Content-Security-Policy "
       "que faz o navegador recusar o script — e não apenas mais uma regra.")
bullet("O CSP estrito (script-src 'self') bloqueia todo script inline, inclusive o "
       "legítimo do próprio DVWA. Em uma aplicação real, o correto seria usar nonces ou "
       "hashes por script, permitindo os legítimos e barrando os injetados.")
bullet("O alvo é o DVWA (deliberadamente vulnerável), conforme o próprio enunciado "
       "permite. As vulnerabilidades da app-api (credencial exposta e command injection) "
       "são intencionais e documentadas, para compor a cadeia de movimento lateral.")

H("Como isto mapeia para a infraestrutura real", 4)
bullet("Firewall nftables no gateway → Security Groups e NACLs na AWS; regras de "
       "roteamento no Proxmox. A camada de rede do provedor precisa permitir o trânsito "
       "até o appliance, que então aplica a política fina — exatamente o que a cadeia "
       "DOCKER-USER + nftables reproduziu localmente.")
bullet("Segmentos internal (DMZ/APP/DB) → subnets/VPCs separadas por camada, sem rota "
       "para a internet, com egress-filtering explícito.")
bullet("WAF Nginx + ModSecurity + CRS → WAF gerenciado (AWS WAF) ou Cloudflare na borda.")
bullet("Gateway WireGuard + SSH por chave via VPN → VPN gerenciada / bastion host com "
       "acesso administrativo isolado.")
bullet("Default-deny em INPUT e FORWARD → postura deny-all com liberações explícitas e "
       "comentadas, princípio de menor privilégio aplicado à rede.")

d.save(OUT)
print("OK ->", OUT)
print("paragrafos:", len(d.paragraphs), "| tabelas:", len(d.tables))
